#!/usr/bin/env python3
"""build_slides.py -- Question 6 presentation for FA24-BEDS-109.

Builds slides/FA24-BEDS-109.pptx with four slides: the variant, the makeover
(default and redesign side by side), the main literacy finding, and the
supporting numbers.  Every number is recomputed from the PSLM files and
cross-checked against scripts/band_table_R.csv; nothing is typed in by hand.

Figures are inserted as SVG pictures with a PNG fallback, the same way
PowerPoint stores Insert > Pictures > .svg.

Run from the project root:  python scripts/build_slides.py
"""

import os
import shutil
import tempfile
import zipfile

import numpy as np
import pandas as pd
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

# --- Registration number and variant -------------------------------------
REG = "FA24-BEDS-109"
VARIANT = int(REG[-2:]) % 8
VARIANTS = [
    ("Punjab", "rural", 2, 1),
    ("Punjab", "urban", 2, 2),
    ("Sindh", "rural", 3, 1),
    ("Sindh", "urban", 3, 2),
    ("Khyber Pakhtunkhwa", "rural", 1, 1),
    ("Khyber Pakhtunkhwa", "urban", 1, 2),
    ("Balochistan", "rural", 4, 1),
    ("Balochistan", "urban", 4, 2),
]
PROV_NAME, AREA_NAME, PROV, URBAN = VARIANTS[VARIANT]
PLACE = f"{PROV_NAME} {AREA_NAME}"

# --- Data ----------------------------------------------------------------
hh = pd.read_csv("data/pslm_hh_subset.csv")
per = pd.read_csv("data/pslm_person_subset.csv")
assert hh["hhid"].is_unique, "hhid must be unique"
assert (hh["hhweight"] > 0).all() and hh["hhweight"].notna().all(), "weights must be positive"
d = per.merge(hh[["hhid", "hhweight", "prov_code", "urban_rural"]], on="hhid", how="inner")
assert len(d) == len(per), "merge changed the person count"
d = d[(d["prov_code"] == PROV) & (d["urban_rural"] == URBAN)].copy()


def wmean(x, w):
    x = np.asarray(x, float); w = np.asarray(w, float)
    return float((x * w).sum() / w.sum())


def wquantile(x, w, q):
    x = np.asarray(x, float); w = np.asarray(w, float)
    o = np.argsort(x, kind="mergesort")
    x, w = x[o], w[o]
    cw = np.cumsum(w) / w.sum()
    return float(x[np.searchsorted(cw, q, side="left")])


# Question 2
students = d[d["ed_total"] > 0]
sp = students["ed_total"].to_numpy(float)
ww = students["hhweight"].to_numpy(float)
logsp = np.log10(sp)
keep = (logsp >= 2) & (logsp <= 6)
sp_k, w_k = sp[keep], ww[keep]
n_students = len(students)
n_analysed = int(keep.sum())
n_excluded = int((~keep).sum())
u_med = wquantile(sp_k, np.ones_like(sp_k), 0.50)
w_med = wquantile(sp_k, w_k, 0.50)
p10 = wquantile(sp_k, w_k, 0.10)
p90 = wquantile(sp_k, w_k, 0.90)
ratio = p90 / p10
mw_below = float(ww[sp < u_med].mean())
mw_above = float(ww[sp > u_med].mean())

# Question 3
ls = d[(d["age"] >= 10) & (d["age"] <= 64) & d["literate"].notna()].copy()
ls["band_i"] = ((ls["age"] - 10) // 5).astype(int)
bands = [f"{10 + 5 * i}-{14 + 5 * i}" for i in range(11)]
rows = []
for i, band in enumerate(bands):
    for sex, sname in ((1, "Men"), (2, "Women")):
        g = ls[(ls["band_i"] == i) & (ls["sex"] == sex)]
        rows.append(dict(band=band, mid=12 + 5 * i, sex=sname, n=len(g),
                         rate=100 * wmean(g["literate"], g["hhweight"]) if len(g) else np.nan))
tab = pd.DataFrame(rows)
men = tab[tab["sex"] == "Men"].reset_index(drop=True)
women = tab[tab["sex"] == "Women"].reset_index(drop=True)
gaps = (men["rate"] - women["rate"]).abs()
gi = int(gaps.idxmax())
gap_band = men.loc[gi, "band"]
gap_value = float(gaps.loc[gi])
gap_mid = float(men.loc[gi, "mid"])
gap_men = float(men.loc[gi, "rate"])
gap_women = float(women.loc[gi, "rate"])

# Cross-check against the R band table written by 03_literacy.R
tab_r = pd.read_csv("scripts/band_table_R.csv")
assert list(tab_r["band"]) == bands, "band labels differ from the R table"
assert (tab_r["men_n"].to_numpy() == men["n"].to_numpy()).all(), "men counts differ"
assert (tab_r["women_n"].to_numpy() == women["n"].to_numpy()).all(), "women counts differ"
assert np.allclose(tab_r["men_rate"], men["rate"], atol=0.05), "men rates differ"
assert np.allclose(tab_r["women_rate"], women["rate"], atol=0.05), "women rates differ"

# Cross-check the Python table and use it for the Q5 claim
tab_py = pd.read_csv("scripts/band_table_python.csv")
py_w19 = float(tab_py.loc[tab_py["band"] == "15-19", "women_rate"].iloc[0])


def rate(band, sex):
    row = tab[(tab["band"] == band) & (tab["sex"] == sex)]
    return float(row["rate"].iloc[0])


w19, m19 = rate("15-19", "Women"), rate("15-19", "Men")
w55, m55 = rate("55-59", "Women"), rate("55-59", "Men")
assert abs(py_w19 - w19) <= 0.05, "Python and R women 15-19 rates differ"
n_literacy = len(ls)
n_cells = len(tab)
min_cell = int(tab["n"].min())
thin_cells = int((tab["n"] < 30).sum())

# --- Formatting helpers ---------------------------------------------------
def num(x):
    return f"{int(round(x)):,}"


def rs(x):
    return f"Rs {int(round(x)):,}"


def pct(x):
    return f"{x:.1f}%"


# --- Theme (midnight analytical) -----------------------------------------
BG = RGBColor(0x0E, 0x1B, 0x2A)
PANEL = RGBColor(0x15, 0x26, 0x3A)
CARD = RGBColor(0xF5, 0xF7, 0xFA)
TEXT = RGBColor(0xF2, 0xF6, 0xFA)
MUTED = RGBColor(0xA9, 0xB7, 0xC6)
TEAL = RGBColor(0x2E, 0xC4, 0xB6)
ORANGE = RGBColor(0xFF, 0x9F, 0x1C)
RED = RGBColor(0xE7, 0x6F, 0x51)
INK = RGBColor(0x10, 0x1E, 0x2C)
BODY = "Segoe UI"

# --- PowerPoint / SVG helpers --------------------------------------------
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def rasterise(svg_path, png_path, scale=1.7):
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPM
    drawing = svg2rlg(svg_path)
    drawing.width *= scale
    drawing.height *= scale
    drawing.scale(scale, scale)
    renderPM.drawToFile(drawing, png_path, fmt="PNG", dpi=150)


def add_svg_picture(slide, svg_path, x, y, w, h, tmpdir):
    """Insert an SVG (with PNG fallback) fitted and centred inside a box."""
    png_path = os.path.join(tmpdir, os.path.splitext(os.path.basename(svg_path))[0] + ".png")
    rasterise(svg_path, png_path)
    with Image.open(png_path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(png_path, Inches(px), Inches(py),
                                   width=Inches(pw), height=Inches(ph))

    # Attach the vector SVG part and reference it from asvg:svgBlip.
    partname = slide.part.package.next_partname("/ppt/media/image%d.svg")
    with open(svg_path, "rb") as fh:
        part = Part(partname, "image/svg+xml", slide.part.package, fh.read())
    svg_rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find("{%s}blip" % A)
    ext_lst = blip.find("{%s}extLst" % A)
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, "{%s}extLst" % A)
    ext = etree.SubElement(ext_lst, "{%s}ext" % A)
    ext.set("uri", SVG_EXT_URI)
    svg_blip = etree.SubElement(ext, "{%s}svgBlip" % ASVG)
    svg_blip.set("{%s}embed" % REL, svg_rid)
    return pic


def fix_svg_content_type(path):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                if 'Extension="svg"' not in text:
                    start = text.index("<Types")
                    close = text.index(">", start) + 1
                    text = (text[:close]
                            + '<Default Extension="svg" ContentType="image/svg+xml"/>'
                            + text[close:])
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def R(t, size=14, color=TEXT, bold=False, italic=False, font=BODY, spc=None):
    spec = dict(t=t, size=size, color=color, bold=bold, italic=italic, font=font)
    if spc is not None:
        spec["spc"] = spc
    return spec


def P(*runs, align=PP_ALIGN.LEFT, before=0, after=0, line=1.0):
    return dict(runs=list(runs), align=align, before=before, after=after, line=line)


def add_text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0)):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0]); tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2]); tf.margin_bottom = Inches(margins[3])
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec["align"]
        p.space_before = Pt(spec["before"])
        p.space_after = Pt(spec["after"])
        p.line_spacing = spec["line"]
        for rspec in spec["runs"]:
            r = p.add_run()
            r.text = rspec["t"]
            f = r.font
            f.name = rspec["font"]
            f.size = Pt(rspec["size"])
            f.bold = rspec["bold"]
            f.italic = rspec["italic"]
            f.color.rgb = rspec["color"]
            if "spc" in rspec:
                r._r.get_or_add_rPr().set("spc", str(rspec["spc"]))
    return tb


def box(slide, x, y, w, h, fill, line=None, rounded=False, radius=0.04, line_w=0.75):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    if rounded:
        shp.adjustments[0] = radius
    return shp


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color


def kicker(slide, x, y, text, color=TEAL, size=11):
    add_text(slide, x, y, 8.0, 0.28, [P(R(text.upper(), size=size, color=color,
                                            bold=True, spc=180))])


def footer(slide, left_text, page):
    add_text(slide, 0.6, 7.08, 10.9, 0.24,
             [P(R(left_text, size=8.5, color=MUTED))])
    add_text(slide, 12.0, 7.08, 0.75, 0.24,
             [P(R(f"{page} / 4", size=8.5, color=MUTED), align=PP_ALIGN.RIGHT)])


# --- Build the deck -------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
tmpdir = tempfile.mkdtemp(prefix="slides109_")

try:
    # ------------------------------------------------------------ Slide 1 --
    s = prs.slides.add_slide(blank)
    set_bg(s, BG)
    box(s, 0.6, 1.25, 0.06, 3.55, TEAL)
    kicker(s, 0.86, 1.18, "ECO346  ·  Assignment 1  ·  Fall 2026", TEAL, 11)
    add_text(s, 0.86, 1.6, 8.2, 1.95, [
        P(R("Literacy, age and", size=42, bold=True), line=0.95),
        P(R("the gender gap", size=42, bold=True), line=0.95),
    ])
    add_text(s, 0.86, 3.78, 7.7, 1.0, [
        P(R(f"A weighted view of {PLACE}, ages 10-64", size=18, color=MUTED)),
        P(R("PSLM 2019-20", size=18, color=MUTED)),
    ])
    add_text(s, 0.86, 6.55, 7.5, 0.3,
             [P(R("Independent practice copy - not the submitted assignment",
                  size=10.5, color=MUTED))])

    box(s, 9.35, 1.2, 3.35, 4.95, PANEL, line=TEAL, rounded=True, radius=0.05)
    add_text(s, 9.7, 1.5, 2.7, 0.3,
             [P(R(f"VARIANT {VARIANT}", size=12, bold=True, color=TEAL, spc=150))])
    add_text(s, 9.7, 1.9, 2.7, 1.2, [
        P(R(PROV_NAME, size=25, bold=True), line=0.95),
        P(R(AREA_NAME, size=25, bold=True), line=0.95),
    ])
    box(s, 9.7, 3.55, 2.6, 0.015, MUTED)
    add_text(s, 9.7, 3.78, 2.7, 0.35,
             [P(R(REG, size=16, bold=True, color=ORANGE))])
    add_text(s, 9.7, 4.45, 2.7, 1.1, [
        P(R("Second language", size=10, color=MUTED, spc=80)),
        P(R("Python", size=14, bold=True)),
        P(R("Weighted by hhweight", size=10, color=MUTED), before=7),
    ])

    # ------------------------------------------------------------ Slide 2 --
    s = prs.slides.add_slide(blank)
    set_bg(s, BG)
    kicker(s, 0.6, 0.45, "The makeover", TEAL)
    add_text(s, 0.6, 0.78, 12.1, 0.85, [
        P(R(f"From {num(n_literacy)} raw points to {n_cells} conditional means",
            size=26, bold=True))])
    add_text(s, 0.6, 1.55, 11.9, 0.42, [
        P(R("The default scatter hides the age pattern under overlapping 0/1 points. "
            "The redesign plots the weighted literacy rate in each band.",
            size=13, color=MUTED))])

    add_text(s, 0.6, 1.95, 5.9, 0.26,
             [P(R("BEFORE  ·  R DEFAULTS", size=10, bold=True, color=RED, spc=120))])
    add_text(s, 6.78, 1.95, 5.9, 0.26,
             [P(R("AFTER  ·  WEIGHTED CONDITIONAL MEANS", size=10, bold=True,
                  color=TEAL, spc=120))])
    box(s, 0.6, 2.25, 5.95, 3.5, CARD, rounded=True, radius=0.03)
    box(s, 6.78, 2.25, 5.95, 3.5, CARD, rounded=True, radius=0.03)
    add_svg_picture(s, "figures/default.svg", 0.78, 2.4, 5.59, 3.2, tmpdir)
    add_svg_picture(s, "figures/literacy.svg", 6.96, 2.4, 5.59, 3.2, tmpdir)

    add_text(s, 0.6, 5.88, 5.95, 0.3,
             [P(R(f"{num(n_literacy)} persons, ages 10-64", size=9.5, color=MUTED))])
    add_text(s, 6.78, 5.88, 5.95, 0.3,
             [P(R(f"{n_cells} band-sex cells; smallest cell n = {num(min_cell)}",
                  size=9.5, color=MUTED))])
    add_text(s, 0.6, 6.35, 12.1, 0.55, [
        P(R("Conditional means remove overplotting", size=12, bold=True, color=TEXT),
          R("      ·      ", size=12, color=TEAL),
          R("Position carries both continuous variables", size=12, bold=True, color=TEXT),
          R("      ·      ", size=12, color=TEAL),
          R("Direct labels replace the legend", size=12, bold=True, color=TEXT))])
    footer(s, "Redesign principles: overplotting, perceptual ranking of channels, data-ink ratio", 2)

    # ------------------------------------------------------------ Slide 3 --
    s = prs.slides.add_slide(blank)
    set_bg(s, BG)
    kicker(s, 0.6, 0.45, "Most important finding", ORANGE)
    add_text(s, 0.6, 0.95, 3.4, 1.25,
             [P(R(f"{gap_value:.1f}", size=78, bold=True, color=ORANGE), line=0.9)])
    add_text(s, 0.65, 2.18, 3.4, 0.3,
             [P(R("POINTS OF DIFFERENCE", size=10, bold=True, color=MUTED, spc=140))])
    add_text(s, 0.6, 2.62, 5.6, 1.5, [
        P(R("The gender gap in literacy is widest at ages ", size=23, bold=True),
          R(gap_band, size=23, bold=True, color=ORANGE)),
    ])
    add_text(s, 0.6, 4.05, 5.7, 0.45, [
        P(R("Men ", size=17, color=TEAL), R(f"{gap_men:.1f}%", size=17, bold=True),
          R("   ·   Women ", size=17, color=MUTED), R(f"{gap_women:.1f}%", size=17, bold=True))])

    box(s, 0.6, 4.95, 2.65, 1.05, PANEL, rounded=True, radius=0.06)
    box(s, 3.42, 4.95, 2.65, 1.05, PANEL, rounded=True, radius=0.06)
    add_text(s, 0.82, 5.12, 2.3, 0.24,
             [P(R("AGES 15-19", size=9, bold=True, color=TEAL, spc=120))])
    add_text(s, 0.82, 5.42, 2.3, 0.4,
             [P(R("W ", size=12, color=MUTED), R(pct(w19), size=15, bold=True),
                R("   M ", size=12, color=MUTED), R(pct(m19), size=15, bold=True))])
    add_text(s, 3.64, 5.12, 2.3, 0.24,
             [P(R("AGES 55-59", size=9, bold=True, color=ORANGE, spc=120))])
    add_text(s, 3.64, 5.42, 2.3, 0.4,
             [P(R("W ", size=12, color=MUTED), R(pct(w55), size=15, bold=True),
                R("   M ", size=12, color=MUTED), R(pct(m55), size=15, bold=True))])

    box(s, 6.55, 1.0, 6.2, 5.55, CARD, rounded=True, radius=0.03)
    add_svg_picture(s, "figures/literacy.svg", 6.75, 1.18, 5.8, 4.55, tmpdir)
    add_text(s, 6.78, 5.8, 5.75, 0.55, [
        P(R("Weighted literacy rate by age band. Hollow markers would flag cells with "
            "fewer than 30 observations; here the smallest cell is ",
            size=8.5, color=INK),
          R(f"n = {num(min_cell)}", size=8.5, bold=True, color=INK),
          R(".", size=8.5, color=INK))])
    footer(s, "Source: PSLM 2019-20 person and household files  ·  weighted by hhweight", 3)

    # ------------------------------------------------------------ Slide 4 --
    s = prs.slides.add_slide(blank)
    set_bg(s, BG)
    kicker(s, 0.6, 0.45, "What the numbers say", TEAL)
    add_text(s, 0.6, 0.78, 11.0, 0.55, [
        P(R("Spending is skewed; the literacy gap is the headline", size=24, bold=True))])

    box(s, 0.6, 1.65, 6.35, 4.95, CARD, rounded=True, radius=0.03)
    add_svg_picture(s, "figures/distribution.svg", 0.78, 1.8, 5.99, 4.05, tmpdir)
    add_text(s, 0.82, 5.92, 5.9, 0.55, [
        P(R("Annual education spending (log10 scale). Weighted median ",
            size=8.5, color=INK),
          R(rs(w_med), size=8.5, bold=True, color=INK),
          R("; 90th percentile ", size=8.5, color=INK),
          R(rs(p90), size=8.5, bold=True, color=INK), R(".", size=8.5, color=INK))])

    box(s, 7.2, 1.65, 5.5, 4.95, PANEL, rounded=True, radius=0.04)
    stat_rows = [
        ("STUDENTS ANALYSED", num(n_analysed), f"{num(n_excluded)} excluded"),
        ("WEIGHTED MEDIAN", rs(w_med), None),
        ("UNWEIGHTED MEDIAN", rs(u_med), None),
        ("P10  ·  P90", f"{rs(p10)}  ·  {rs(p90)}", None),
        ("P90 / P10", f"{ratio:.2f}x", None),
        ("WOMEN 15-19", pct(w19), f"R = Python ({pct(py_w19)})"),
    ]
    yy = 1.92
    for label, value, note in stat_rows:
        add_text(s, 7.5, yy, 4.9, 0.22,
                 [P(R(label, size=8.5, bold=True, color=MUTED, spc=110))])
        add_text(s, 7.5, yy + 0.24, 4.9, 0.38,
                 [P(R(value, size=16, bold=True, color=TEXT))])
        if note:
            add_text(s, 7.5, yy + 0.57, 4.9, 0.2,
                     [P(R(note, size=8.5, color=MUTED))])
            yy += 0.84
        else:
            yy += 0.64
        if label != stat_rows[-1][0]:
            box(s, 7.5, yy - 0.16, 4.85, 0.01, RGBColor(0x2A, 0x3E, 0x55))
    footer(s, "Single cross-section: cannot separate cohort effects from ageing  ·  R and Python agree", 4)

    # ------------------------------------------------------------- Save ---
    os.makedirs("slides", exist_ok=True)
    out = f"slides/{REG}.pptx"
    prs.save(out)
    fix_svg_content_type(out)
finally:
    shutil.rmtree(tmpdir, ignore_errors=True)

print(f"Saved {out} with {len(prs.slides)} slides")
print(f"Variant {VARIANT}: {PLACE} (prov_code {PROV}, urban_rural {URBAN})")
print(f"Q2: students={n_students}, analysed={n_analysed}, excluded={n_excluded}, "
      f"w_med={rs(w_med)}, u_med={rs(u_med)}, p10={rs(p10)}, p90={rs(p90)}, ratio={ratio:.2f}")
print(f"Q3: women15-19={w19:.1f}, men15-19={m19:.1f}, women55-59={w55:.1f}, "
      f"men55-59={m55:.1f}, gap={gap_band} {gap_value:.1f}, min_cell={min_cell}, thin={thin_cells}")
print(f"Q5: women15-19 R={w19:.1f}, Python={py_w19:.1f}")
